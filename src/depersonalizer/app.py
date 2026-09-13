# -*- coding: utf-8 -*-
"""Графический интерфейс приложения (tkinter)."""

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from docx import Document

from .excel_core import (
    check_key_compatible,
    depersonalize,
    read_key,
    read_table,
    restore,
    write_plain_table,
    _rows_to_text,
)
from .word_core import (
    CATEGORY_ORDER,
    depersonalize_docx,
    restore_docx,
    _iter_docx_paragraphs,
)

FILE_TYPES = [("Excel-файлы", "*.xlsx"), ("Все файлы", "*.*")]
WORD_FILE_TYPES = [("Word-файлы", "*.docx"), ("Все файлы", "*.*")]


class ProgressWindow:
    def __init__(self, parent, title, maximum):
        self.win = tk.Toplevel(parent)
        self.win.title(title)
        self.win.resizable(False, False)
        tk.Label(self.win, text="Выполняется конвертация…").pack(padx=20, pady=(15, 10))
        self.bar = ttk.Progressbar(self.win, mode="determinate", maximum=maximum, length=320)
        self.bar.pack(padx=20, pady=(0, 15))
        self.win.update_idletasks()
        parent.winfo_toplevel().wait_window_flag = False  # noqa

    def step(self):
        self.bar.step(1)
        self.win.update_idletasks()

    def close(self):
        self.win.destroy()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Деперсонализатор")
        self.geometry("520x380")
        self.resizable(False, False)

        tk.Label(
            self,
            text="Деперсонализатор данных",
            font=("Segoe UI", 16, "bold"),
        ).pack(pady=(25, 5))
        tk.Label(
            self,
            text="Локальная обработка xlsx и Word. Данные не покидают ваш компьютер.",
            fg="#555",
        ).pack(pady=(0, 15))

        tk.Label(self, text="Excel (xlsx):", fg="#777").pack(pady=(0, 2))
        tk.Button(
            self,
            text="Прямая конвертация (деперсонализация)",
            width=42,
            command=self.open_depersonalize_window,
        ).pack(pady=4)
        tk.Button(
            self,
            text="Обратная конвертация (восстановление)",
            width=42,
            command=self.open_restore_window,
        ).pack(pady=4)

        tk.Label(self, text="Word (docx):", fg="#777").pack(pady=(8, 2))
        tk.Button(
            self,
            text="Прямая конвертация (Word)",
            width=42,
            command=self.open_depersonalize_word_window,
        ).pack(pady=4)
        tk.Button(
            self,
            text="Обратная конвертация (Word)",
            width=42,
            command=self.open_restore_word_window,
        ).pack(pady=4)

        tk.Button(
            self,
            text="Выход",
            width=42,
            command=self.destroy,
        ).pack(pady=(8, 0))

        self._dep_window = None
        self._restore_window = None
        self._dep_word_window = None
        self._restore_word_window = None

    # -------------------------------------------------------------- windows

    def open_depersonalize_window(self):
        path = filedialog.askopenfilename(
            title="Выберите исходный файл xlsx",
            filetypes=FILE_TYPES,
        )
        if not path:
            return
        try:
            headers, rows = read_table(path)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось прочитать файл:\n{e}")
            return
        if not headers:
            messagebox.showerror("Ошибка", "Файл пуст: в нём нет ни одной колонки.")
            return

        if self._dep_window and self._dep_window.winfo_exists():
            self._dep_window.destroy()
        win = tk.Toplevel(self)
        self._dep_window = win
        win.title("Выбор колонок")
        win.geometry("460x420")
        win.resizable(False, False)

        tk.Label(win, text=f"Файл: {path}", wraplength=420, justify="left").pack(
            padx=12, pady=(12, 4)
        )
        tk.Label(win, text="Отметьте колонки, значения которых заменить на индексы:").pack(
            anchor="w", padx=12
        )

        canvas = tk.Canvas(win, borderwidth=0, highlightthickness=0)
        scroll = ttk.Scrollbar(win, orient="vertical", command=canvas.yview)
        inner = ttk.Frame(canvas)
        inner.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all")),
        )
        canvas.create_window((0, 0), window=inner, anchor="nw")
        canvas.configure(yscrollcommand=scroll.set)
        canvas.pack(side="left", fill="both", expand=True, padx=(12, 0), pady=8)
        scroll.pack(side="right", fill="y", pady=8, padx=(0, 12))

        check_vars = {}
        for i, header in enumerate(headers):
            var = tk.BooleanVar(value=False)
            check_vars[i] = var
            ttk.Checkbutton(inner, text=f"{i + 1}. {header}", variable=var).pack(
                anchor="w"
            )

        tk.Label(
            win,
            text="Пустые ячейки не заменяются. Одинаковые значения получают один индекс.",
            fg="#555",
            wraplength=420,
            justify="left",
        ).pack(padx=12, pady=4)

        def on_confirm():
            selected = {i for i, v in check_vars.items() if v.get()}
            if not selected:
                messagebox.showerror("Ошибка", "Отметьте хотя бы одну колонку.")
                return
            names = ", ".join(f"{i + 1}. {headers[i]}" for i in sorted(selected))
            if not messagebox.askyesno(
                "Подтверждение",
                "Будут заменены на индексы колонки:" + "\n" + names,
            ):
                return
            self.run_depersonalize(win, path, headers, rows, selected)

        tk.Button(win, text="Деперсонализировать", command=on_confirm).pack(
            pady=(4, 12)
        )

    def open_restore_window(self):
        if self._restore_window and self._restore_window.winfo_exists():
            self._restore_window.destroy()
        win = tk.Toplevel(self)
        self._restore_window = win
        win.title("Обратная конвертация")
        win.geometry("480x280")
        win.resizable(False, False)

        state = {"data_path": "", "key_path": ""}

        def pick_data():
            state["data_path"] = filedialog.askopenfilename(
                title="Выберите деперсонализированный файл",
                filetypes=FILE_TYPES,
            )
            data_lbl.config(text=state["data_path"] or "Файл не выбран")

        def pick_key():
            state["key_path"] = filedialog.askopenfilename(
                title="Выберите файл ключа",
                filetypes=FILE_TYPES,
            )
            key_lbl.config(text=state["key_path"] or "Файл не выбран")

        tk.Label(win, text="Деперсонализированный файл:").pack(anchor="w", padx=12, pady=(15, 2))
        tk.Button(win, text="Выбрать файл", command=pick_data).pack(anchor="w", padx=12)
        data_lbl = tk.Label(win, text="Файл не выбран", fg="#555", wraplength=440, justify="left")
        data_lbl.pack(anchor="w", padx=12)

        tk.Label(win, text="Файл ключа:").pack(anchor="w", padx=12, pady=(10, 2))
        tk.Button(win, text="Выбрать ключ", command=pick_key).pack(anchor="w", padx=12)
        key_lbl = tk.Label(win, text="Файл не выбран", fg="#555", wraplength=440, justify="left")
        key_lbl.pack(anchor="w", padx=12)

        def on_restore():
            if not state["data_path"]:
                messagebox.showerror("Ошибка", "Не выбран деперсонализированный файл.")
                return
            if not state["key_path"]:
                messagebox.showerror("Ошибка", "Не выбран файл ключа.")
                return
            self.run_restore(win, state["data_path"], state["key_path"])

        tk.Button(win, text="Восстановить", command=on_restore).pack(pady=(18, 12))

    # ------------------------------------------------------------ pipelines

    def run_depersonalize(self, win, src_path, headers, rows, selected):
        out_path = filedialog.asksaveasfilename(
            parent=win,
            title="Куда сохранить деперсонализированный файл?",
            defaultextension=".xlsx",
            initialfile=src_path.rsplit("/", 1)[-1].rsplit("\\", 1)[-1].replace(
                ".xlsx", "_деперсонализированные.xlsx"
            ),
            filetypes=FILE_TYPES,
        )
        if not out_path:
            return
        key_path = filedialog.asksaveasfilename(
            parent=win,
            title="Куда сохранить ключевую таблицу?",
            defaultextension=".xlsx",
            initialfile=out_path.replace("_деперсонализированные.xlsx", "_ключ.xlsx"),
            filetypes=FILE_TYPES,
        )
        if not key_path:
            return

        self.withdraw()
        progress = ProgressWindow(win, "Деперсонализация", maximum=max(1, len(rows) + 2))

        try:
            new_rows, key_rows, _ = depersonalize(headers, rows, selected)
            progress.step()
            write_plain_table(out_path, headers, _rows_to_text(new_rows))
            progress.step()
            write_plain_table(
                key_path,
                ["Исходная_колонка", "Индекс", "Исходное_значение"],
                key_rows,
            )
            progress.step()
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать файлы:\n{e}")
            return
        finally:
            progress.close()
            self.deiconify()

        messagebox.showinfo(
            "Готово",
            f"Файлы сохранены:\n{out_path}\n\nКлюч:\n{key_path}\n\n"
            "Храните ключ отдельно от данных — без него восстановление невозможно.",
        )

    def run_restore(self, win, data_path, key_path):
        try:
            headers, rows = read_table(data_path)
            key_headers, mapping = read_key(key_path)
            check_key_compatible(headers, key_headers)
        except Exception as e:
            messagebox.showerror("Ошибка", str(e))
            return

        out_path = filedialog.asksaveasfilename(
            parent=win,
            title="Куда сохранить восстановленный файл?",
            defaultextension=".xlsx",
            initialfile=data_path.rsplit("/", 1)[-1].rsplit("\\", 1)[-1].replace(
                ".xlsx", "_восстановлено.xlsx"
            ),
            filetypes=FILE_TYPES,
        )
        if not out_path:
            return
        if out_path == data_path and not messagebox.askyesno(
            "Внимание",
            "Путь сохранения совпадает с исходным файлом. Файл будет перезаписан. Продолжить?",
        ):
            return

        self.withdraw()
        progress = ProgressWindow(win, "Восстановление", maximum=max(1, len(rows) + 2))
        try:
            new_rows = restore(headers, rows, mapping)
            progress.step()
            write_plain_table(out_path, headers, _rows_to_text(new_rows))
            progress.step()
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать файл:\n{e}")
            return
        finally:
            progress.close()
            self.deiconify()

        messagebox.showinfo("Готово", f"Файл сохранён:\n{out_path}")

    # -------------------------------------------------- Word pipelines

    def open_depersonalize_word_window(self):
        path = filedialog.askopenfilename(
            title="Выберите исходный файл Word (docx)",
            filetypes=WORD_FILE_TYPES,
        )
        if not path:
            return
        try:
            Document(path)
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось прочитать файл:\n{e}")
            return

        if self._dep_word_window and self._dep_word_window.winfo_exists():
            self._dep_word_window.destroy()
        win = tk.Toplevel(self)
        self._dep_word_window = win
        win.title("Выбор категорий")
        win.geometry("480x640")
        win.resizable(True, True)

        tk.Label(win, text=f"Файл: {path}", wraplength=440, justify="left").pack(
            padx=12, pady=(12, 4)
        )
        tk.Label(win, text="Отметьте категории данных для замены на индексы:").pack(
            anchor="w", padx=12
        )

        cat_vars = {}
        for cat in CATEGORY_ORDER:
            var = tk.BooleanVar(value=False)
            cat_vars[cat] = var
            ttk.Checkbutton(win, text=cat, variable=var).pack(anchor="w", padx=24)

        tk.Label(
            win,
            text="Дополнительные значения (по одному на строку):",
        ).pack(anchor="w", padx=12, pady=(10, 2))
        txt = tk.Text(win, height=5, width=50)
        txt.pack(padx=12, pady=(0, 4))

        tk.Label(
            win,
            text="Одинаковые значения получают один индекс. "
            "Форматирование текста рядом с заменённым значением может не сохраниться.",
            fg="#555",
            wraplength=420,
            justify="left",
        ).pack(padx=12, pady=4)

        def on_confirm():
            categories = [c for c in CATEGORY_ORDER if cat_vars[c].get()]
            custom_values = [
                line.strip() for line in txt.get("1.0", "end").splitlines() if line.strip()
            ]
            if not categories and not custom_values:
                messagebox.showerror(
                    "Ошибка",
                    "Отметьте хотя бы одну категорию или введите свои значения.",
                )
                return
            message = "Будут заменены следующие данные:" + "\n"
            message += "Категории: " + ("; ".join(categories) if categories else "не выбраны")
            if custom_values:
                message += f"\nДополнительные значения: {len(custom_values)} шт."
            if not messagebox.askyesno("Подтверждение", message):
                return
            self.run_depersonalize_word(win, path, categories, custom_values)

        tk.Button(win, text="Деперсонализировать", command=on_confirm).pack(
            pady=(4, 12)
        )

    def open_restore_word_window(self):
        if self._restore_word_window and self._restore_word_window.winfo_exists():
            self._restore_word_window.destroy()
        win = tk.Toplevel(self)
        self._restore_word_window = win
        win.title("Обратная конвертация (Word)")
        win.geometry("480x280")
        win.resizable(False, False)

        state = {"data_path": "", "key_path": ""}

        def pick_data():
            state["data_path"] = filedialog.askopenfilename(
                title="Выберите деперсонализированный файл Word",
                filetypes=WORD_FILE_TYPES,
            )
            data_lbl.config(text=state["data_path"] or "Файл не выбран")

        def pick_key():
            state["key_path"] = filedialog.askopenfilename(
                title="Выберите файл ключа",
                filetypes=FILE_TYPES,
            )
            key_lbl.config(text=state["key_path"] or "Файл не выбран")

        tk.Label(win, text="Деперсонализированный файл (docx):").pack(
            anchor="w", padx=12, pady=(15, 2)
        )
        tk.Button(win, text="Выбрать файл", command=pick_data).pack(anchor="w", padx=12)
        data_lbl = tk.Label(
            win, text="Файл не выбран", fg="#555", wraplength=440, justify="left"
        )
        data_lbl.pack(anchor="w", padx=12)

        tk.Label(win, text="Файл ключа:").pack(anchor="w", padx=12, pady=(10, 2))
        tk.Button(win, text="Выбрать ключ", command=pick_key).pack(anchor="w", padx=12)
        key_lbl = tk.Label(
            win, text="Файл не выбран", fg="#555", wraplength=440, justify="left"
        )
        key_lbl.pack(anchor="w", padx=12)

        def on_restore():
            if not state["data_path"]:
                messagebox.showerror("Ошибка", "Не выбран деперсонализированный файл.")
                return
            if not state["key_path"]:
                messagebox.showerror("Ошибка", "Не выбран файл ключа.")
                return
            self.run_restore_word(win, state["data_path"], state["key_path"])

        tk.Button(win, text="Восстановить", command=on_restore).pack(pady=(18, 12))

    def run_depersonalize_word(self, win, src_path, categories, custom_values):
        out_path = filedialog.asksaveasfilename(
            parent=win,
            title="Куда сохранить деперсонализированный документ?",
            defaultextension=".docx",
            initialfile=src_path.rsplit("/", 1)[-1].rsplit("\\", 1)[-1].replace(
                ".docx", "_деперсонализированные.docx"
            ),
            filetypes=WORD_FILE_TYPES,
        )
        if not out_path:
            return
        key_path = filedialog.asksaveasfilename(
            parent=win,
            title="Куда сохранить ключевую таблицу?",
            defaultextension=".xlsx",
            initialfile=out_path.replace(
                "_деперсонализированные.docx", "_ключ.docx.xlsx"
            ),
            filetypes=FILE_TYPES,
        )
        if not key_path:
            return

        try:
            tmp = Document(src_path)
            total = sum(1 for _ in _iter_docx_paragraphs(tmp))
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось прочитать файл:\n{e}")
            return

        self.withdraw()
        progress = ProgressWindow(win, "Деперсонализация Word", maximum=max(1, total + 2))
        try:
            key_rows, _ = depersonalize_docx(
                src_path, categories, custom_values, progress, out_path
            )
            progress.step()
            write_plain_table(
                key_path,
                ["Категория", "Индекс", "Исходное_значение"],
                key_rows,
            )
            progress.step()
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать файлы:\n{e}")
            return
        finally:
            progress.close()
            self.deiconify()

        messagebox.showinfo(
            "Готово",
            f"Файлы сохранены:\n{out_path}\n\nКлюч:\n{key_path}\n\n"
            "Храните ключ отдельно от данных — без него восстановление невозможно.",
        )

    def run_restore_word(self, win, data_path, key_path):
        out_path = filedialog.asksaveasfilename(
            parent=win,
            title="Куда сохранить восстановленный документ?",
            defaultextension=".docx",
            initialfile=data_path.rsplit("/", 1)[-1].rsplit("\\", 1)[-1].replace(
                ".docx", "_восстановлено.docx"
            ),
            filetypes=WORD_FILE_TYPES,
        )
        if not out_path:
            return
        if out_path == data_path and not messagebox.askyesno(
            "Внимание",
            "Путь сохранения совпадает с исходным файлом. Файл будет перезаписан. Продолжить?",
        ):
            return

        try:
            tmp = Document(data_path)
            total = sum(1 for _ in _iter_docx_paragraphs(tmp))
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось прочитать файл:\n{e}")
            return

        self.withdraw()
        progress = ProgressWindow(win, "Восстановление Word", maximum=max(1, total + 2))
        try:
            restore_docx(data_path, key_path, progress, out_path)
            progress.step()
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось создать файл:\n{e}")
            return
        finally:
            progress.close()
            self.deiconify()

        messagebox.showinfo("Готово", f"Файл сохранён:\n{out_path}")


def main():
    app = App()
    app.mainloop()