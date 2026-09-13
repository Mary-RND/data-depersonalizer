# -*- coding: utf-8 -*-
"""Работа с Word (.docx): маскирование категорий и восстановление.

Категории распознаются по встроенным правилам (регулярным выражениям),
одинаковые значения в пределах категории получают один индекс.
"""

import re

from docx import Document

from .excel_core import _cell_text

WORD_PATTERNS = {
    "Телефоны": re.compile(
        r"(?<![0-9])(?:\+7|8)\s*[(\- ]?\s*\d{1,3}[\s)\-]*\d{1,3}[\s\-]*\d{2}[\s\-]*\d{2}(?!\d)",
        re.IGNORECASE,
    ),
    "Email": re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}"),
    "ФИО": re.compile(
        r"(?<![А-ЯЁA-Z]\.)(?<!\.)(?<![\w])[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+)?",
    ),
    "Адреса": re.compile(
        r"(?<!\w)(?:\d{6}[,\s]+\s*)?"
        r"(?:"
        r"[А-ЯЁ][а-яё-]+(?:\s+[А-ЯЁ][а-яё-]+)?\s+"
        r"(?:[Рр]еспублик\w*|[Оо]бласт\w*|[Оо]бл\.?|[Кк]рай)[,\s]+|"
        r"[Рр]еспублик\w*\s+[А-ЯЁ][а-яё-]+[,\s]+"
        r")?"
        r"(?:"
        r"(?:г(?:ород|\.))[ ,]?\s*[А-ЯЁ][а-яё-]+|"
        r"пос(?:елок)?\.?\s+[А-ЯЁа-яё0-9-]+|"
        r"пгт\s+[А-ЯЁа-яё0-9-]+|"
        r"село\s+[А-ЯЁа-яё0-9-]+|"
        r"(?:с\.|дер(?:евня)?\.?)\s+[А-ЯЁа-яё0-9-]+|"
        r"мкр(?:он)?\.?\s+[А-ЯЁа-яё0-9-]+|"
        r"терр(?:итория)?\.?\s+[А-ЯЁа-яё0-9-]+|"
        r"[А-ЯЁ][а-яё-]+(?:ская|цкая)\s+область"
        r")"
        r"(?:[,\s]{1,3}(?:"
        r"ул\.\s+[А-ЯЁа-яё0-9-]+|"
        r"улиц(?:а|ы)\s+[А-ЯЁа-яё0-9-]+|"
        r"пр-т\s+[А-ЯЁа-яё0-9-]+|"
        r"проспект\s+[А-ЯЁа-яё0-9-]+|"
        r"пер\.\s+[А-ЯЁа-яё0-9-]+|"
        r"переул(?:ок|ка)\s+[А-ЯЁа-яё0-9-]+|"
        r"наб\.\s+[А-ЯЁа-яё0-9-]+|"
        r"набережн(?:ая|ой)\s+[А-ЯЁа-яё0-9-]+|"
        r"(?:бульвар|шоссе|проезд|аллея)\s+[А-ЯЁа-яё0-9-]+|"
        r"(?:д\.|дом|стр\.|строение|корп\.|корпус|кв\.|квартир[аы]|офис|пом\.|помещение|этаж)"
        r"\s*№?\s*[0-9А-ЯЁа-яё()/\-\.]+|"
        r"[А-ЯЁ][а-яё-]+(?:,\s?\d{1,4})?"
        r")){0,7}"
    ),
    "Паспортные данные": re.compile(
        r"(?<![\d])(?:[Сс]ерия\s*[:\s]\s*)?(?:\d{4}|\d{2}\s\d{2})\s(?:№\s*)?\d{6}(?!\d)"
        r"|(?<![А-Яа-яЁё0-9])код\s*подразделения\s*\d{3}\s*[-\s]\s*\d{3}"
    ),
    "Названия компаний": re.compile(
        r"(?<![\w«»\"'])(?:"
        r"(?:ООО|ОАО|ЗАО|АО|ПАО|ИП|НКО|ТСЖ|ТОО)\s*[«»\"'][^«»\"'\r\n;]{2,70}[«»\"']|"
        r"(?:ООО|ОАО|ЗАО|АО|ПАО|НКО)\s+[А-ЯЁ][А-ЯЁа-яёA-Za-z0-9&()\-.]{1,40}"
        r"(?:\s+[А-ЯЁ]?[А-ЯЁа-яёA-Za-z0-9&()\-.]{1,40}){0,3}"
        r")(?=\s*[,;.()]|$|\s+[А-ЯЁ0-9A-Z«»\"'])"
    ),
    "Реквизиты организации": re.compile(
        r"(?<![А-Яа-яЁё0-9])(?:ИНН|КПП|ОГРНИП|ОГРН|ОКПО|ОКВЭД|ОКТМО|ОКАТО|СНИЛС)"
        r"\s*[:\s—\-. ]?\s*\d{9,15}"
        r"|(?<!\d)\d{3}[-\s]\d{3}[-\s]\d{3}\s\d{2}(?!\d)"
    ),
    "Банковские реквизиты": re.compile(
        r"(?<![А-Яа-яЁё0-9])(?:"
        r"р/с(?:ч)?|к/с(?:ч)?|б/с|л/с(?:ч)?|корсчет|кор/счет|корр\w+|БИК|"
        r"расчетн\w*\s*[сc]чет\w*|расчётн\w*\s*[сc]чёт\w*|"
        r"корреспондентск\w*\s*(?:[сc]чет|счёт)\w*|лицев\w*\s*[сc]чет\w*"
        r")\s*[:\s№\-. ]?\s*\d{9,25}"
    ),
}

PHONE_DIGITS_RE = re.compile(r"\d+")


def _normalize_phone(text):
    return "".join(PHONE_DIGITS_RE.findall(text))


def _phone_token(original, mapping, counter):
    digits = _normalize_phone(original)
    if len(digits) < 10:
        return original, mapping, counter
    key = ("phone", digits)
    if key in mapping:
        return mapping[key], mapping, counter
    counter += 1
    token = f"ИДТЕЛ_{counter:04d}"
    mapping[key] = token
    mapping[(token,)] = original
    return token, mapping, counter


def _email_token(original, mapping, counter):
    key = ("email", original.lower())
    if key in mapping:
        return mapping[key], mapping, counter
    counter += 1
    token = f"ИДЕМЕЙЛ_{counter:04d}"
    mapping[key] = token
    mapping[(token,)] = original
    return token, mapping, counter


def _fio_token(original, mapping, counter):
    key = ("fio", original)
    if key in mapping:
        return mapping[key], mapping, counter
    counter += 1
    token = f"ИДФИО_{counter:04d}"
    mapping[key] = token
    mapping[(token,)] = original
    return token, mapping, counter


def _make_generic_replacer(kind, prefix):
    """Создаёт токенизатор для категории с точным сравнением исходного значения."""

    def tokenizer(original, mapping, counter):
        key = (kind, original)
        if key in mapping:
            return mapping[key], mapping, counter
        counter += 1
        token = f"{prefix}_{counter:04d}"
        mapping[key] = token
        mapping[(token,)] = original
        return token, mapping, counter

    return tokenizer


CATEGORY_TOKENIZERS = {
    "Телефоны": _phone_token,
    "Email": _email_token,
    "ФИО": _fio_token,
    "Адреса": _make_generic_replacer("address", "ИДАДРЕС"),
    "Паспортные данные": _make_generic_replacer("passport", "ИДПАСП"),
    "Названия компаний": _make_generic_replacer("company", "ИДКОМП"),
    "Реквизиты организации": _make_generic_replacer("req", "ИДРЕКВ"),
    "Банковские реквизиты": _make_generic_replacer("bank", "ИДБАНК"),
}
CATEGORY_PATTERNS = {c: WORD_PATTERNS[c] for c in CATEGORY_TOKENIZERS}
# Порядок замены важен: короткие числовые значения заменяются раньше
# длинных (адресов, компаний, ФИО), чтобы токены не попадали внутрь замен.
CATEGORY_ORDER = [
    "Телефоны",
    "Email",
    "Паспортные данные",
    "Реквизиты организации",
    "Банковские реквизиты",
    "Адреса",
    "Названия компаний",
    "ФИО",
]


def _replace_custom_values(text, custom_values, mapping, counters):
    """Заменяет произвольные пользовательские значения на индексы ИДДОП_<NNNN>."""
    for value in custom_values:
        if not value or value not in text:
            continue
        key = ("custom", value)
        if key in mapping:
            token = mapping[key]
        else:
            counters["Дополнительно"] = counters.get("Дополнительно", 0) + 1
            token = f"ИДДОП_{counters['Дополнительно']:04d}"
            mapping[key] = token
            mapping[(token,)] = value
        text = text.replace(value, token)
    return text


def _replace_in_text(text, categories, mapping, counters, custom_values):
    """Заменяет найденные персональные данные в тексте на индексные токены."""
    text = _replace_custom_values(text, custom_values, mapping, counters)
    for cat in CATEGORY_ORDER:
        if cat in categories:
            text = _replace_category(text, cat, mapping, counters)
    return text


def _replace_category(text, cat, mapping, counters):
    def repl(match):
        original = match.group(0)
        counter = counters.get(cat, 0)
        token, _, counter = CATEGORY_TOKENIZERS[cat](
            original, mapping, counter
        )
        counters[cat] = counter
        return token

    return CATEGORY_PATTERNS[cat].sub(repl, text)


def _iter_docx_paragraphs(doc):
    """Перебирает все абзацы: тело документа, заголовки/колонтитулы, ячейки таблиц."""
    for p in doc.paragraphs:
        yield p
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            yield p
        for p in sec.footer.paragraphs:
            yield p
        for tbl in sec.header.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        yield p
        for tbl in sec.footer.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        yield p

    def walk_tables(tables):
        for tbl in tables:
            for row in tbl.rows:
                for cell in row.cells:
                    yield cell

    def walk_cells(tables):
        for cell in walk_tables(tables):
            for p in cell.paragraphs:
                yield p
            nested = cell.tables
            if nested:
                for p in walk_cells(nested):
                    yield p

    for p in walk_cells(doc.tables):
        yield p


def _set_paragraph_text(p, text):
    """Заменяет весь текст абзаца одним прогоном (упрощённая плоская замена)."""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def depersonalize_docx(path, categories, custom_values=None, progress=None, out_path=None):
    """Деперсонализирует .docx. Возвращает (key_rows, replaced_count).

    key_rows - строки для ключевой таблицы
               (Категория, Индекс, Исходное_значение).
    Если задан out_path - деперсонализированный документ сохраняется туда.
    """
    custom_values = [v.strip() for v in (custom_values or []) if v and v.strip()]
    doc = Document(path)
    mapping = {}
    counters = {c: 0 for c in categories}

    for p in _iter_docx_paragraphs(doc):
        text = p.text
        if not text:
            continue
        new_text = _replace_in_text(text, categories, mapping, counters, custom_values)
        if new_text != text:
            _set_paragraph_text(p, new_text)
        if progress:
            progress.step()

    key_rows = []
    seen = {}
    ordered_kinds = list(_category_kinds(categories))
    if custom_values:
        ordered_kinds.append("custom")
    for token_kind in ordered_kinds:
        for key, token in mapping.items():
            if not isinstance(key, tuple) or len(key) != 2:
                continue
            kind, _ = key
            if kind != token_kind:
                continue
            value_key = (token,)
            if value_key not in mapping:
                continue
            original = mapping[value_key]
            pair = (kind, original)
            if pair not in seen:
                seen[pair] = token
                key_rows.append([CATEGORY_BY_KIND.get(kind, "Дополнительно"), token, original])

    if out_path:
        doc.save(out_path)

    replaced_count = len(key_rows)
    return key_rows, replaced_count


CATEGORY_BY_KIND = {
    "phone": "Телефоны",
    "email": "Email",
    "fio": "ФИО",
    "address": "Адреса",
    "passport": "Паспортные данные",
    "company": "Названия компаний",
    "req": "Реквизиты организации",
    "bank": "Банковские реквизиты",
}

KIND_BY_CATEGORY = {
    "Телефоны": "phone",
    "Email": "email",
    "ФИО": "fio",
    "Адреса": "address",
    "Паспортные данные": "passport",
    "Названия компаний": "company",
    "Реквизиты организации": "req",
    "Банковские реквизиты": "bank",
}


def _category_kinds(categories):
    return [KIND_BY_CATEGORY[c] for c in CATEGORY_ORDER if c in categories]


def read_key_word(path):
    """Читает ключ для Word. Возвращает mapping {категория: {индекс: значение}}."""
    headers, rows = read_table(path)
    if not headers:
        return {}
    cat_idx = headers.index("Категория") if "Категория" in headers else -1
    index_idx = headers.index("Индекс") if "Индекс" in headers else -1
    value_idx = headers.index("Исходное_значение") if "Исходное_значение" in headers else -1
    if cat_idx < 0 or index_idx < 0 or value_idx < 0:
        raise ValueError("Файл ключа имеет неверную структуру: "
                         "ожидаются колонки «Категория», «Индекс», «Исходное_значение».")
    mapping = {}
    for row in rows:
        cat = _cell_text(row[cat_idx])
        index = _cell_text(row[index_idx])
        if not cat or not index:
            continue
        mapping.setdefault(cat, {})[index] = row[value_idx]
    return mapping


def restore_docx_text(text, mapping):
    """Заменяет индексные токены на исходные значения."""
    for cat_map in mapping.values():
        for index, value in cat_map.items():
            if index and index in text:
                text = text.replace(index, _cell_text(value))
    return text


def restore_docx(path, key_path, progress=None, out_path=None):
    """Восстанавливает исходные значения в .docx по ключу.

    Если задан out_path - восстановленный документ сохраняется туда.
    """
    mapping = read_key_word(key_path)
    if not mapping:
        raise ValueError("Файл ключа пуст: не найдено ни одного соответствия.")
    doc = Document(path)
    for p in _iter_docx_paragraphs(doc):
        text = p.text
        if not text:
            continue
        new_text = restore_docx_text(text, mapping)
        if new_text != text:
            _set_paragraph_text(p, new_text)
        if progress:
            progress.step()
    if out_path:
        doc.save(out_path)
    return mapping