# -*- coding: utf-8 -*-
"""Работа с Excel (xlsx): прямая и обратная конвертация.

Модуль без зависимостей от графического интерфейса:
все функции принимают и возвращают обычные данные.
"""

from openpyxl import Workbook, load_workbook


def _cell_text(value):
    """Возвращает строковое представление значения ячейки или пустую строку."""
    if value is None:
        return ""
    return str(value).strip()


def read_table(path):
    """Читает первый лист xlsx. Возвращает (headers, rows).

    headers - список строк заголовков (первая строка листа),
    rows     - список строк с данными (все строки кроме первой).
    """
    wb = load_workbook(path, read_only=False, data_only=False)
    ws = wb.worksheets[0]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not all_rows:
        return [], []

    headers = []
    for i, value in enumerate(all_rows[0]):
        if value is None or _cell_text(value) == "":
            headers.append(f"Колонка {i + 1}")
        else:
            headers.append(_cell_text(value))
    rows = [list(r) for r in all_rows[1:]]
    return headers, rows


def _rows_to_text(rows):
    return [[_cell_text(v) for v in row] for row in rows]


def depersonalize(headers, rows, selected_positions):
    """Заменяет значения в выбранных колонках на индексы.

    selected_positions - набор 0-индексных позиций колонок.
    Возвращает (new_rows, key_rows, replaced_count).
    key_rows - строки для ключевой таблицы
               (Исходная_колонка, Индекс, Исходное_значение).
    """
    new_rows = [list(row) for row in rows]
    key_rows = []

    for pos in sorted(selected_positions):
        col_no = pos + 1  # номер колонки слева направо (1-based)
        header = headers[pos]
        mapping = {}
        counter = 0
        for row in new_rows:
            value = row[pos]
            if value is None:
                continue
            text = _cell_text(value)
            if text == "":
                continue
            if value not in mapping:
                counter += 1
                index = f"ИД{col_no}_{counter:04d}"
                mapping[value] = index
                key_rows.append([header, index, value])
            row[pos] = mapping[value]

    replaced_count = len(key_rows)
    return new_rows, key_rows, replaced_count


def write_plain_table(path, headers, rows):
    """Записывает таблицу (заголовки + данные) в xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Лист 1"
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    wb.close()


def read_key(path):
    """Читает ключевой файл. Возвращает (header_values, mapping).

    mapping - {заголовок колонки: {индекс: исходное значение}}.
    header_values - множество заголовков из ключа.
    """
    headers, rows = read_table(path)
    if not headers:
        return set(), {}
    name_idx = headers.index("Исходная_колонка") if "Исходная_колонка" in headers else -1
    index_idx = headers.index("Индекс") if "Индекс" in headers else -1
    value_idx = headers.index("Исходное_значение") if "Исходное_значение" in headers else -1
    if name_idx < 0 or index_idx < 0 or value_idx < 0:
        raise ValueError("Файл ключа имеет неверную структуру: "
                         "ожидаются колонки «Исходная_колонка», «Индекс», «Исходное_значение».")

    mapping = {}
    header_values = set()
    for row in rows:
        header = _cell_text(row[name_idx])
        if not header:
            continue
        header_values.add(header)
        index = _cell_text(row[index_idx])
        value = row[value_idx]
        mapping.setdefault(header, {})[index] = value
    return header_values, mapping


def restore(headers, rows, mapping):
    """Восстанавливает исходные значения по ключу.

    Возвращает новые строки. Ячейки, чьё значение совпадает с индексом
    из ключа для своей колонки, заменяются исходным значением.
    """
    new_rows = []
    for row in rows:
        new_row = list(row)
        for pos, value in enumerate(new_row):
            if pos >= len(headers) or value is None:
                continue
            header = headers[pos]
            col_map = mapping.get(header)
            if not col_map:
                continue
            text = _cell_text(value)
            if text in col_map:
                new_row[pos] = col_map[text]
        new_rows.append(new_row)
    return new_rows


def check_key_compatible(headers, key_headers):
    """Проверка совместимости ключа и файла по структуре."""
    if not key_headers:
        raise ValueError("Файл ключа пуст: не найдено ни одной колонки соответствий.")
    file_set = set(_cell_text(h) for h in headers)
    if not (file_set & key_headers):
        raise ValueError("Ключ не совпадает с файлом: ни одна колонка ключа "
                         "не найдена в обрабатываемом файле.")
    return True